#!/usr/bin/env python3
"""Analyze multiple Phyphox CSV trials for Lab 03 Plans 1, 2, and 3.

For each raw CSV, the program detects time/acceleration columns, calculates
velocity and position with trapezoidal integration, and makes graphs. It also
creates one Excel workbook and a summary CSV of the final results.

The program also averages the trial results within each plan. Because trials
may have different sample times, it interpolates them onto a common elapsed-
time grid before calculating and graphing the average curves.

Interactive: python3 phyphox_motion_analyzer.py
Command line: python3 phyphox_motion_analyzer.py --plan1 p1.csv --plan2 p2.csv --plan3 p3.csv
"""

from __future__ import annotations

import argparse
import re
import shlex
import sys
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd


@dataclass
class TrialResult:
    plan: int
    trial: int
    source_file: Path
    axis: str
    acceleration_bias: float
    data: pd.DataFrame
    graph_file: Path
    duration: float
    samples: int
    max_acceleration: float
    min_acceleration: float
    max_velocity: float
    min_velocity: float
    final_velocity: float
    displacement: float


def normalized(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name).lower())


def find_time_column(columns: list[str]) -> str:
    for column in columns:
        name = normalized(column)
        if name in {"time", "times", "timestamp", "t"} or name.startswith("time"):
            return column
    raise ValueError("No time column found. Expected a name such as 'Time (s)'.")


def acceleration_columns(columns: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for column in columns:
        name = normalized(column)
        if "acc" not in name:
            continue
        for axis in "xyz":
            if (re.search(rf"(^|acceleration|acc){axis}($|ms|mss|m)", name)
                    or name.startswith(axis + "acc")):
                result.setdefault(axis, column)
    for column in columns:
        name = normalized(column)
        if name in {"x", "y", "z"}:
            result.setdefault(name, column)
    return result


def choose_axis(data: pd.DataFrame, available: dict[str, str], requested: str) -> str:
    if requested != "auto":
        if requested not in available:
            found = ", ".join(sorted(available)) or "none"
            raise ValueError(f"Axis '{requested}' not found; available axes: {found}.")
        return requested
    if not available:
        raise ValueError("No x, y, or z acceleration columns found.")
    return max(available, key=lambda a: pd.to_numeric(data[available[a]], errors="coerce").std())


def cumulative_trapezoid(values: np.ndarray, time: np.ndarray, initial: float) -> np.ndarray:
    result = np.empty(len(values), dtype=float)
    result[0] = initial
    result[1:] = initial + np.cumsum(0.5 * (values[:-1] + values[1:]) * np.diff(time))
    return result


def read_raw_csv(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, sep=None, engine="python")
    except Exception as error:
        raise ValueError(f"Could not read {path.name}: {error}") from error


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", text).strip("_") or "trial"


def create_graph(data: pd.DataFrame, result_label: str, axis: str,
                 output_path: Path, show: bool) -> None:
    time = data["Time (s)"].to_numpy()
    series = [
        (data[f"Acceleration {axis} (m/s^2)"], "Acceleration", "Acceleration (m/s²)", "#2563eb"),
        (data[f"Velocity {axis} (m/s)"], "Calculated velocity", "Velocity (m/s)", "#16a34a"),
        (data[f"Position {axis} (m)"], "Calculated position", "Position (m)", "#dc2626"),
    ]
    fig, axes = plt.subplots(3, 1, figsize=(10, 10), sharex=True)
    for plot_axis, (values, title, ylabel, color) in zip(axes, series):
        plot_axis.plot(time, values, color=color, linewidth=1.4)
        plot_axis.axhline(0, color="black", linewidth=0.7, alpha=0.5)
        plot_axis.set_title(f"{title} in {axis}-direction")
        plot_axis.set_ylabel(ylabel)
        plot_axis.grid(True, alpha=0.3)
    axes[-1].set_xlabel("Time (s)")
    fig.suptitle(result_label, fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(output_path, dpi=200, bbox_inches="tight")
    if show:
        plt.show()
    plt.close(fig)


def analyze_trial(source_file: Path, plan: int, trial: int,
                  args: argparse.Namespace, graphs_directory: Path) -> TrialResult:
    raw = read_raw_csv(source_file)
    if raw.empty:
        raise ValueError(f"{source_file.name} contains no data.")
    columns = [str(column) for column in raw.columns]
    time_column = find_time_column(columns)
    available = acceleration_columns(columns)
    axis = choose_axis(raw, available, args.axis)
    acceleration_column = available[axis]
    acceleration_name = f"Acceleration {axis} (m/s^2)"

    data = pd.DataFrame({
        "Time (s)": pd.to_numeric(raw[time_column], errors="coerce"),
        acceleration_name: pd.to_numeric(raw[acceleration_column], errors="coerce"),
    }).dropna().sort_values("Time (s)").drop_duplicates("Time (s)").reset_index(drop=True)
    if len(data) < 2:
        raise ValueError(f"{source_file.name} needs at least two valid measurements.")

    time = data["Time (s)"].to_numpy(dtype=float)
    acceleration = data[acceleration_name].to_numpy(dtype=float)
    if np.any(np.diff(time) <= 0):
        raise ValueError(f"Time values in {source_file.name} must increase.")

    acceleration_bias = 0.0
    if args.remove_bias is not None:
        mask = time <= time[0] + args.remove_bias
        if mask.sum() < 2:
            raise ValueError(f"Bias interval in {source_file.name} has fewer than two measurements.")
        acceleration_bias = float(np.mean(acceleration[mask]))
        acceleration = acceleration - acceleration_bias
        data[acceleration_name] = acceleration

    velocity = cumulative_trapezoid(acceleration, time, args.initial_velocity)
    position = cumulative_trapezoid(velocity, time, args.initial_position)
    data["Delta t (s)"] = np.insert(np.diff(time), 0, 0.0)
    data[f"Velocity {axis} (m/s)"] = velocity
    data[f"Position {axis} (m)"] = position

    graph_file = graphs_directory / f"Plan_{plan}_Trial_{trial}_{safe_name(source_file.stem)}.png"
    label = f"Plan {plan}, Trial {trial} — {source_file.name}"
    create_graph(data, label, axis, graph_file, args.show)

    return TrialResult(plan, trial, source_file, axis, acceleration_bias, data, graph_file,
        float(time[-1] - time[0]), len(data), float(np.max(acceleration)),
        float(np.min(acceleration)), float(np.max(velocity)), float(np.min(velocity)),
        float(velocity[-1]), float(position[-1] - position[0]))


def summary_dataframe(results: list[TrialResult]) -> pd.DataFrame:
    rows = []
    for r in results:
        rows.append({
            "Plan": r.plan,
            "Trial": r.trial,
            "Source CSV": r.source_file.name,
            "Source path": str(r.source_file.resolve()),
            "Axis": r.axis,
            "Samples": r.samples,
            "Duration (s)": r.duration,
            "Acceleration bias removed (m/s^2)": r.acceleration_bias,
            "Maximum acceleration (m/s^2)": r.max_acceleration,
            "Minimum acceleration (m/s^2)": r.min_acceleration,
            "Maximum velocity (m/s)": r.max_velocity,
            "Minimum velocity (m/s)": r.min_velocity,
            "Final velocity (m/s)": r.final_velocity,
            "Final displacement (m)": r.displacement,
            "Expected displacement": "Nonzero" if r.plan in {1, 2} else "Approximately zero",
            "Graph file": r.graph_file.name,
        })
    return pd.DataFrame(rows)


def plan_average_curve(plan_results: list[TrialResult]) -> pd.DataFrame:
    """Return point-by-point average curves over the trials' shared duration."""
    if not plan_results:
        raise ValueError("At least one trial is required to calculate a plan average.")

    shared_duration = min(r.duration for r in plan_results)
    point_count = min(r.samples for r in plan_results)
    common_time = np.linspace(0.0, shared_duration, point_count)
    acceleration_curves = []
    velocity_curves = []
    position_curves = []

    for r in plan_results:
        elapsed_time = r.data["Time (s)"].to_numpy(dtype=float)
        elapsed_time = elapsed_time - elapsed_time[0]
        acceleration_curves.append(np.interp(
            common_time, elapsed_time,
            r.data[f"Acceleration {r.axis} (m/s^2)"].to_numpy(dtype=float)))
        velocity_curves.append(np.interp(
            common_time, elapsed_time,
            r.data[f"Velocity {r.axis} (m/s)"].to_numpy(dtype=float)))
        position = r.data[f"Position {r.axis} (m)"].to_numpy(dtype=float)
        position_curves.append(np.interp(common_time, elapsed_time, position - position[0]))

    acceleration_array = np.vstack(acceleration_curves)
    velocity_array = np.vstack(velocity_curves)
    position_array = np.vstack(position_curves)
    return pd.DataFrame({
        "Elapsed time (s)": common_time,
        "Average acceleration (m/s^2)": acceleration_array.mean(axis=0),
        "Acceleration standard deviation (m/s^2)": acceleration_array.std(axis=0),
        "Average velocity (m/s)": velocity_array.mean(axis=0),
        "Velocity standard deviation (m/s)": velocity_array.std(axis=0),
        "Average position (m)": position_array.mean(axis=0),
        "Position standard deviation (m)": position_array.std(axis=0),
    })


def create_average_graph(data: pd.DataFrame, plan: int, trial_count: int,
                         output_path: Path, show: bool) -> None:
    time = data["Elapsed time (s)"].to_numpy()
    series = [
        ("Average acceleration (m/s^2)", "Acceleration standard deviation (m/s^2)",
         "Average acceleration", "Acceleration (m/s²)", "#2563eb"),
        ("Average velocity (m/s)", "Velocity standard deviation (m/s)",
         "Average velocity", "Velocity (m/s)", "#16a34a"),
        ("Average position (m)", "Position standard deviation (m)",
         "Average position", "Position (m)", "#dc2626"),
    ]
    fig, axes = plt.subplots(3, 1, figsize=(10, 10), sharex=True)
    for plot_axis, (mean_column, std_column, title, ylabel, color) in zip(axes, series):
        mean = data[mean_column].to_numpy()
        standard_deviation = data[std_column].to_numpy()
        plot_axis.plot(time, mean, color=color, linewidth=2, label="Trial average")
        plot_axis.fill_between(time, mean - standard_deviation, mean + standard_deviation,
                               color=color, alpha=0.18, label="±1 standard deviation")
        plot_axis.axhline(0, color="black", linewidth=0.7, alpha=0.5)
        plot_axis.set_title(title)
        plot_axis.set_ylabel(ylabel)
        plot_axis.grid(True, alpha=0.3)
        plot_axis.legend(loc="best")
    axes[-1].set_xlabel("Elapsed time (s)")
    fig.suptitle(f"Plan {plan} Average of {trial_count} Trial(s)",
                 fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(output_path, dpi=200, bbox_inches="tight")
    if show:
        plt.show()
    plt.close(fig)


def calculate_plan_averages(results: list[TrialResult], graphs_directory: Path,
                            output_directory: Path, show: bool
                            ) -> tuple[pd.DataFrame, dict[int, pd.DataFrame]]:
    rows = []
    curves_by_plan: dict[int, pd.DataFrame] = {}
    for plan in sorted({r.plan for r in results}):
        plan_results = [r for r in results if r.plan == plan]
        curve = plan_average_curve(plan_results)
        graph_file = graphs_directory / f"Plan_{plan}_Average_Curves.png"
        curve_csv = output_directory / f"Plan_{plan}_Average_Curves.csv"
        curve.to_csv(curve_csv, index=False)
        create_average_graph(curve, plan, len(plan_results), graph_file, show)
        curves_by_plan[plan] = curve

        rows.append({
            "Plan": plan,
            "Number of trials": len(plan_results),
            "Average samples": float(np.mean([r.samples for r in plan_results])),
            "Average duration (s)": float(np.mean([r.duration for r in plan_results])),
            "Average acceleration bias removed (m/s^2)": float(np.mean([r.acceleration_bias for r in plan_results])),
            "Average maximum acceleration (m/s^2)": float(np.mean([r.max_acceleration for r in plan_results])),
            "Average minimum acceleration (m/s^2)": float(np.mean([r.min_acceleration for r in plan_results])),
            "Average maximum velocity (m/s)": float(np.mean([r.max_velocity for r in plan_results])),
            "Average minimum velocity (m/s)": float(np.mean([r.min_velocity for r in plan_results])),
            "Average final velocity (m/s)": float(np.mean([r.final_velocity for r in plan_results])),
            "Average final displacement (m)": float(np.mean([r.displacement for r in plan_results])),
            "Expected displacement": "Nonzero" if plan in {1, 2} else "Approximately zero",
            "Average curve CSV": curve_csv.name,
            "Average graph file": graph_file.name,
        })
    return pd.DataFrame(rows), curves_by_plan


def write_excel_workbook(results: list[TrialResult], summary: pd.DataFrame,
                         plan_averages: pd.DataFrame,
                         average_curves: dict[int, pd.DataFrame],
                         workbook_path: Path) -> None:
    try:
        import openpyxl  # noqa: F401
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError("Install the Excel dependency with: python3 -m pip install openpyxl") from error

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        plan_averages.to_excel(writer, sheet_name="Plan Averages", index=False)
        sheet_results: list[tuple[str, TrialResult]] = []
        for r in results:
            sheet_name = f"P{r.plan}_Trial_{r.trial}"[:31]
            r.data.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet_results.append((sheet_name, r))
        for plan, curve in average_curves.items():
            curve.to_excel(writer, sheet_name=f"P{plan}_Average_Curves", index=False)

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            for column_cells in worksheet.columns:
                longest = max(len(str(cell.value or "")) for cell in column_cells)
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(longest + 2, 12), 42)
        workbook["Summary"].sheet_view.showGridLines = False
        workbook["Plan Averages"].sheet_view.showGridLines = False

        for sheet_name, r in sheet_results:
            worksheet = workbook[sheet_name]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            value_headers = [f"Acceleration {r.axis} (m/s^2)", f"Velocity {r.axis} (m/s)", f"Position {r.axis} (m)"]
            titles = ["Acceleration vs Time", "Velocity vs Time", "Position vs Time"]
            anchors = ["G2", "G18", "G34"]
            for header, title, anchor in zip(value_headers, titles, anchors):
                chart = LineChart()
                chart.title = title
                chart.y_axis.title = header
                chart.x_axis.title = "Time (s)"
                chart.height, chart.width = 7.2, 14.5
                values = Reference(worksheet, min_col=headers[header], min_row=1, max_row=worksheet.max_row)
                times = Reference(worksheet, min_col=headers["Time (s)"], min_row=2, max_row=worksheet.max_row)
                chart.add_data(values, titles_from_data=True)
                chart.set_categories(times)
                chart.legend = None
                worksheet.add_chart(chart, anchor)

        for plan in average_curves:
            worksheet = workbook[f"P{plan}_Average_Curves"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            average_headers = [
                "Average acceleration (m/s^2)",
                "Average velocity (m/s)",
                "Average position (m)",
            ]
            titles = ["Average Acceleration vs Time", "Average Velocity vs Time",
                      "Average Position vs Time"]
            anchors = ["I2", "I18", "I34"]
            for header, title, anchor in zip(average_headers, titles, anchors):
                chart = LineChart()
                chart.title = title
                chart.y_axis.title = header
                chart.x_axis.title = "Elapsed time (s)"
                chart.height, chart.width = 7.2, 14.5
                values = Reference(worksheet, min_col=headers[header], min_row=1,
                                   max_row=worksheet.max_row)
                times = Reference(worksheet, min_col=headers["Elapsed time (s)"], min_row=2,
                                  max_row=worksheet.max_row)
                chart.add_data(values, titles_from_data=True)
                chart.set_categories(times)
                chart.legend = None
                worksheet.add_chart(chart, anchor)


def parse_path_line(entered: str) -> Path | None:
    try:
        parts = shlex.split(entered)
    except ValueError:
        return None
    return Path(parts[0]).expanduser() if len(parts) == 1 else None


def ask_for_plan_files(plan: int) -> list[Path]:
    print(f"\nPLAN {plan}: enter or drag one raw CSV at a time.")
    print("Press Enter on a blank line when this plan is complete.")
    files: list[Path] = []
    while True:
        entered = input(f"Plan {plan} CSV #{len(files) + 1}: ").strip()
        if not entered:
            return files
        if entered.lower() in {"q", "quit", "exit"}:
            raise KeyboardInterrupt
        path = parse_path_line(entered)
        if path is None:
            print("Please enter one CSV file at a time.")
        elif not path.is_file():
            print(f"File not found: {path}")
        elif path.suffix.lower() != ".csv":
            print("That file is not a CSV file.")
        else:
            files.append(path)
            print(f"Added: {path.name}")


def collect_interactive_files() -> dict[int, list[Path]]:
    print("Phyphox Multi-Plan Motion Analyzer")
    print("Add raw data files for Plans 1, 2, and 3.")
    return {plan: ask_for_plan_files(plan) for plan in (1, 2, 3)}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Analyze multiple Phyphox CSV files for Lab 03.")
    parser.add_argument("--plan1", type=Path, nargs="*", default=[], help="Plan 1 CSV files")
    parser.add_argument("--plan2", type=Path, nargs="*", default=[], help="Plan 2 CSV files")
    parser.add_argument("--plan3", type=Path, nargs="*", default=[], help="Plan 3 CSV files")
    parser.add_argument("--axis", choices=["auto", "x", "y", "z"], default="auto")
    parser.add_argument("--initial-velocity", type=float, default=0.0)
    parser.add_argument("--initial-position", type=float, default=0.0)
    parser.add_argument("--remove-bias", type=float, metavar="SECONDS")
    parser.add_argument("--output-directory", type=Path, default=Path("Lab03_Analysis_Output"))
    parser.add_argument("--show", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    plan_files = {1: args.plan1, 2: args.plan2, 3: args.plan3}
    if not any(plan_files.values()):
        try:
            plan_files = collect_interactive_files()
        except (EOFError, KeyboardInterrupt):
            print("\nProgram ended without processing files.")
            return 0
    if not any(plan_files.values()):
        print("No CSV files selected.", file=sys.stderr)
        return 1
    for files in plan_files.values():
        for path in files:
            if not path.is_file():
                print(f"Error: file not found: {path}", file=sys.stderr)
                return 1

    output_directory = args.output_directory.expanduser()
    graphs_directory = output_directory / "graphs"
    output_directory.mkdir(parents=True, exist_ok=True)
    graphs_directory.mkdir(parents=True, exist_ok=True)

    results: list[TrialResult] = []
    failures: list[str] = []
    for plan in (1, 2, 3):
        for trial, source_file in enumerate(plan_files[plan], start=1):
            print(f"Analyzing Plan {plan}, Trial {trial}: {source_file.name}")
            try:
                result = analyze_trial(source_file, plan, trial, args, graphs_directory)
                results.append(result)
                print(f"  axis={result.axis}, final velocity={result.final_velocity:.4f} m/s, displacement={result.displacement:.4f} m")
            except (ValueError, OSError) as error:
                failures.append(f"Plan {plan}, {source_file.name}: {error}")
                print(f"  Skipped: {error}", file=sys.stderr)
    if not results:
        print("No files could be analyzed.", file=sys.stderr)
        return 1

    summary = summary_dataframe(results)
    summary_path = output_directory / "Lab03_Final_Results.csv"
    averages_path = output_directory / "Lab03_Plan_Averages.csv"
    workbook_path = output_directory / "Lab03_Motion_Analysis.xlsx"
    summary.to_csv(summary_path, index=False)
    plan_averages, average_curves = calculate_plan_averages(
        results, graphs_directory, output_directory, args.show)
    plan_averages.to_csv(averages_path, index=False)
    try:
        write_excel_workbook(results, summary, plan_averages, average_curves, workbook_path)
    except RuntimeError as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1

    print("\nAnalysis complete.")
    print(f"Excel workbook: {workbook_path.resolve()}")
    print(f"Final-results CSV: {summary_path.resolve()}")
    print(f"Plan-averages CSV: {averages_path.resolve()}")
    print(f"Graph folder: {graphs_directory.resolve()}")
    if failures:
        print("\nSkipped files:")
        for failure in failures:
            print(f"  - {failure}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
